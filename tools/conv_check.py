#!/usr/bin/env python3
"""Check Parade's converters with other software.

  python3 tools/conv_check.py BUILD_DIR

Expects BUILD_DIR/conv/rich.* (written by test_convert with PARADE_CONV_OUT)
and BUILD_DIR/sample.jdoc (from pd_dump). Every external program runs under
an address-space cap and a timeout.
"""
import os
import resource
import shutil
import subprocess
import sys
import tempfile
import zipfile

BUILD = sys.argv[1] if len(sys.argv) > 1 else "build"
CONV = os.path.join(BUILD, "conv")
PD_CONV = os.path.join(BUILD, "pd_conv")
MEM_MB = int(os.environ.get("CONV_CHECK_MEM_MB", "4096"))
fails = 0
checks = 0


def cap():
    resource.setrlimit(resource.RLIMIT_AS, (MEM_MB << 20, MEM_MB << 20))


def run(cmd, timeout=180, cwd=None):
    return subprocess.run(cmd, cwd=cwd, preexec_fn=cap, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                          timeout=timeout)


def check(cond, what):
    global fails, checks
    checks += 1
    print(("  ok    " if cond else "  FAIL  ") + what)
    if not cond:
        fails += 1


KEY = ["Conversion Test", "First Section", "Item one", "Nested item", "A quoted paragraph", "alpha", "1.5",
       "A footnote here", "The note text", "Last paragraph"]


def pdf_text(pdf):
    r = run(["pdftotext", "-layout", pdf, "-"])
    return r.stdout.decode("utf-8", "replace")


def soffice(src, fmt, outdir):
    """convert with LibreOffice into a fresh directory (so no older file can pass for its output)"""
    profile = tempfile.mkdtemp(prefix="lo-profile-")
    fresh = tempfile.mkdtemp(prefix="lo-out-", dir=outdir)
    try:
        run(["soffice", "-env:UserInstallation=file://" + profile, "--headless", "--convert-to", fmt,
             "--outdir", fresh, src], timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    base = os.path.splitext(os.path.basename(src))[0]
    out = os.path.join(fresh, base + "." + fmt.split(":")[0])
    return out if os.path.exists(out) else None


def foreign_docx(work):
    """A document written by python-docx (Word's default template) read by Parade."""
    import docx
    from docx.shared import Inches
    png = os.path.join(work, "dot.png")
    import base64
    open(png, "wb").write(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVQI12P4z8AAAAMBAQAY3Y2wAAAAAElFTkSuQmCC"))
    d = docx.Document()
    d.add_heading("Foreign Title", 0)
    d.add_heading("Section One", level=1)
    p = d.add_paragraph("Plain ")
    p.add_run("bold").bold = True
    p.add_run(" and ")
    p.add_run("italic").italic = True
    d.add_paragraph("first bullet", style="List Bullet")
    d.add_paragraph("second bullet", style="List Bullet")
    d.add_paragraph("first number", style="List Number")
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    for r in range(2):
        for c in range(2):
            t.cell(r, c).text = "cell %d%d" % (r, c)
    d.add_picture(png, width=Inches(1))
    d.add_page_break()
    d.add_paragraph("after the break")
    src = os.path.join(work, "foreign.docx")
    d.save(src)
    pdoc = os.path.join(work, "foreign.jdoc")
    r = run([PD_CONV, src, pdoc])
    check(r.returncode == 0, "pd_conv imports it")
    if r.returncode != 0:
        return
    import json
    j = json.load(open(pdoc, encoding="utf-8"))
    flat = json.dumps(j)
    check('"Role": "title"' in flat, "Title style -> title")
    check('"Role": "heading"' in flat and '"Level": 1' in flat, "Heading 1 -> heading")
    check('"Weight": 700' in flat and '"Italic": true' in flat, "bold and italic runs")
    check(flat.count('"List":') >= 3, "List Bullet / List Number paragraphs are list items")
    check('"_TreeNode_(table)"' in flat and "cell 11" in flat, "the table")
    check('"image/png"' in flat, "the picture")
    check('"Break": "page"' in flat or '"_TreeNode_(break)"' in flat, "the page break")
    r = run([PD_CONV, src, os.path.join(work, "foreign.pdf")])
    check(r.returncode == 0, "and lays it out to PDF")


def main():
    os.makedirs(CONV, exist_ok=True)
    work = tempfile.mkdtemp(prefix="conv-check-")

    print("HTML")
    import lxml.html
    tree = lxml.html.parse(os.path.join(CONV, "rich.html"))
    root = tree.getroot()
    check(len(root.findall(".//h1")) == 2 and len(root.findall(".//h2")) == 1, "headings parse as h1/h2")
    check(len(root.findall(".//table//th")) == 3 and len(root.findall(".//td[@colspan='2']")) == 1,
          "table with header cells and a column span")
    check(len(root.findall(".//figure/img")) + len(root.findall(".//figure//img")) >= 1, "figure with image")
    check(root.find(".//section[@class='footnotes']") is not None, "footnotes section")

    print("Markdown")
    import mistune
    md = mistune.create_markdown(plugins=["strikethrough", "table", "footnotes"])
    html = md(open(os.path.join(CONV, "rich.md"), encoding="utf-8").read())
    for frag in ["<strong>bold</strong>", "<em>italic</em>", "<del>struck</del>", "<table>", "<blockquote>",
                 "<ol>", "<code>code_span</code>", 'href="https://example.com/a?b=1&amp;c=2"']:
        check(frag in html, "mistune renders " + frag)

    print("LaTeX")
    if shutil.which("lualatex"):
        tex = os.path.join(work, "rich.tex")
        shutil.copy(os.path.join(CONV, "rich.tex"), tex)
        ok = True
        for _ in range(2):  # references settle in the second run
            r = run(["lualatex", "-interaction=nonstopmode", "-halt-on-error", "rich.tex"], timeout=300, cwd=work)
            ok = ok and r.returncode == 0
        check(ok and os.path.exists(os.path.join(work, "rich.pdf")), "LuaLaTeX compiles the export")
        if ok:
            t = pdf_text(os.path.join(work, "rich.pdf"))
            check(all(k in t for k in KEY if k != "The note text") and "The note text" in t,
                  "the PDF has all the text")
    else:
        print("  (no lualatex)")

    print("DOCX")
    import docx
    dd = docx.Document(os.path.join(CONV, "rich.docx"))
    texts = [p.text for p in dd.paragraphs]
    check(any("Conversion Test" in t for t in texts), "python-docx reads the paragraphs")
    check(len(dd.tables) == 1 and len(dd.tables[0].rows) == 3, "python-docx sees the table")
    check(any(p.style.name == "heading 1" or p.style.name == "Heading 1" for p in dd.paragraphs),
          "heading style survives")
    with zipfile.ZipFile(os.path.join(CONV, "rich.docx")) as z:
        check(z.testzip() is None, "the zip is valid (CRCs)")
        check("word/footnotes.xml" in z.namelist() and "word/media/image1.png" in z.namelist(),
              "footnotes and media parts")

    print("Parade reads python-docx output")
    foreign_docx(work)

    writer = False
    if shutil.which("soffice"):     # Writer may be missing (Draw alone opens DOCX but not RTF or HTML)
        probe = os.path.join(work, "probe.rtf")
        open(probe, "w").write("{\\rtf1\\ansi{\\fonttbl{\\f0 Times;}}\\pard Hello\\par}")
        writer = soffice(probe, "pdf", work) is not None
        if not writer:
            print("  (LibreOffice Writer is not installed: the LibreOffice checks are skipped)")

    if writer:
        for fmt in ["docx", "rtf", "html"]:
            print("LibreOffice reads " + fmt.upper())
            pdf = soffice(os.path.join(CONV, "rich." + fmt), "pdf", work)
            check(pdf is not None, "converts to PDF")
            if pdf:
                t = pdf_text(pdf)
                missing = [k for k in KEY if k not in t]
                check(not missing, "all key text present" + (" (missing %s)" % missing if missing else ""))
                shutil.copy(pdf, os.path.join(work, "lo-" + fmt + ".pdf"))

        print("Parade reads LibreOffice's output")
        for fmt, filt in [("docx", "docx:MS Word 2007 XML"), ("rtf", "rtf")]:
            src = os.path.join(work, "lo-src.html")
            shutil.copy(os.path.join(CONV, "rich.html"), src)
            out = soffice(src, filt, work)
            check(out is not None, "LibreOffice writes " + fmt)
            if out:
                txt = os.path.join(work, "back-" + fmt + ".txt")
                r = run([PD_CONV, out, txt])
                check(r.returncode == 0, "pd_conv imports it")
                if r.returncode == 0:
                    t = open(txt, encoding="utf-8").read()
                    missing = [k for k in KEY if k not in t]
                    check(not missing, "all key text back" + (" (missing %s)" % missing if missing else ""))
                    pdoc = os.path.join(work, "back-" + fmt + ".jdoc")
                    run([PD_CONV, out, pdoc])
                    s = open(pdoc, encoding="utf-8").read()
                    check('"Role": "heading"' in s, "headings recognized")
                    check('"_TreeNode_(table)"' in s, "the table recognized")
                    check('"footnote"' in s, "the footnote recognized")
                    check('"Weight": 700' in s and '"Italic": true' in s, "bold and italic recognized")
                os.remove(out)


    print("sample document through every format to PDF")
    sample = os.path.join(BUILD, "sample.jdoc")
    for ext in ["html", "md", "rtf", "docx"]:
        mid = os.path.join(work, "sample." + ext)
        r1 = run([PD_CONV, sample, mid])
        r2 = run([PD_CONV, mid, os.path.join(work, "sample-" + ext + ".pdf")])
        check(r1.returncode == 0 and r2.returncode == 0, "sample -> %s -> PDF" % ext)

    shutil.rmtree(work, ignore_errors=True)
    print("%d checks, %d failures" % (checks, fails))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
